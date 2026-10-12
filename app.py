"""
SERVER FLASK & REST API — MASTER SUITE UNIFICADA (Challenger Automotive 3-in-1)
=============================================================================
Servidor Web Flask con endpoints REST API, motor de sincronización automática
en segundo plano y servidor de frontend Tailwind CSS.
"""

from flask import Flask, render_template, jsonify, request, send_file
import os, json, tempfile, uuid, gzip
from datetime import datetime, date

from odoo_sync_engine import (
    CRON_ENGINE, EMPRESAS,
    obtener_vendedores_odoo, calcular_comisiones_engine,
    obtener_productos_engine, obtener_stock_engine, obtener_cartera_cobranza_engine,
    obtener_historial_compras_cliente_engine,
    enviar_correo_cobranza_odoo,
    generar_datos_demo_stock, fmt_mxn, fmt_num
)
from sat_sync_engine import (
    procesar_facturas_sat_puras, procesar_cuentas_por_pagar,
    descargar_cfdis_sat, cargar_config_fiel, guardar_config_fiel,
    obtener_xml_por_uuid, generar_zip_facturas_sat, construir_indice_sat_xmls, parsear_xml_cfdi
)
from ml_forecasting_engine import obtener_historico_top20_ml
from estado_resultados_engine import procesar_estado_resultados

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
app = Flask(
    __name__,
    template_folder=os.path.join(BASE_DIR, "templates"),
    static_folder=os.path.join(BASE_DIR, "static")
)
app.config['TEMPLATES_AUTO_RELOAD'] = True
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 0

import time

# Iniciar el motor de sincronización en segundo plano al arrancar el servidor
CRON_ENGINE.iniciar()

@app.after_request
def add_header(response):
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"

    # Compresión GZIP en vivo para respuestas JSON/texto superiores a 1KB (acelera red hasta 92%)
    accept_encoding = request.headers.get("Accept-Encoding", "")
    if "gzip" in accept_encoding.lower() and response.status_code == 200 and not response.direct_passthrough:
        try:
            data = response.get_data()
            if len(data) > 1024:
                compressed = gzip.compress(data, compresslevel=5)
                response.set_data(compressed)
                response.headers["Content-Encoding"] = "gzip"
                response.headers["Content-Length"] = len(compressed)
        except Exception:
            pass
    return response

@app.route("/")
def home():
    return render_template("index.html", cache_id=int(time.time()))


@app.route("/api/status")
def get_status():
    return jsonify(CRON_ENGINE.get_status())


@app.route("/api/sync_now", methods=["POST"])
def sync_now():
    def _run():
        try:
            CRON_ENGINE.sincronizar_todo()
        except Exception:
            pass
    import threading
    threading.Thread(target=_run, daemon=True).start()
    return jsonify({"success": True, "message": "Sincronización manual iniciada en segundo plano."})


@app.route("/api/vendedores")
def get_vendedores():
    empresa = request.args.get("empresa", "BALDOMERO MEDINA GARZA")
    cid = EMPRESAS.get(empresa, {}).get("company_id", 1) or 1
    try:
        vends = obtener_vendedores_odoo(cid)
        return jsonify({"success": True, "vendedores": vends})
    except Exception as e:
        return jsonify({"success": False, "error": str(e), "vendedores": []})


@app.route("/api/comisiones", methods=["POST"])
def post_comisiones():
    data = request.get_json() or {}
    mes = int(data.get("mes", datetime.now().month))
    anio = int(data.get("anio", datetime.now().year))
    vendedor_val = data.get("vendedor_valor")
    vendedor_nom = data.get("vendedor_nombre", "")
    empresa = data.get("empresa_nombre", "BALDOMERO MEDINA GARZA")
    incluir_pbnk = data.get("incluir_pbnk", True)
    excluir_45 = data.get("excluir_retrasos_45", True)
    limite_dias = int(data.get("limite_dias", 45))
    overrides = data.get("overrides_especiales", {})

    if not vendedor_val:
        return jsonify({"success": False, "error": "Selecciona un vendedor válido."})

    try:
        res = calcular_comisiones_engine(
            mes, anio, vendedor_val, vendedor_nom, empresa,
            overrides_especiales=overrides, incluir_pbnk=incluir_pbnk,
            excluir_retrasos_45=excluir_45, limite_dias=limite_dias
        )
        return jsonify({"success": True, "data": res})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})


HISTORIAL_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "historial_comisiones.json")

def cargar_historial_comisiones():
    if os.path.exists(HISTORIAL_FILE):
        try:
            with open(HISTORIAL_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {}

def guardar_historial_comisiones(data):
    with open(HISTORIAL_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

@app.route("/api/comisiones/pagos", methods=["GET", "POST"])
def api_comisiones_pagos():
    if request.method == "POST":
        data = request.get_json() or {}
        clave = data.get("clave", "").strip()
        com = float(data.get("comision_total", 0))
        monto = float(data.get("monto", 0))
        fecha = data.get("fecha", "").strip()
        nota = data.get("nota", "").strip()

        if not clave or not monto or not fecha:
            return jsonify({"success": False, "error": "Faltan campos requeridos (clave, monto, fecha)"})

        hist = cargar_historial_comisiones()
        if clave not in hist:
            hist[clave] = {"comision_total": com, "pagos": []}
        if com > 0:
            hist[clave]["comision_total"] = com
        hist[clave]["pagos"].append({"fecha": fecha, "monto": monto, "nota": nota})
        guardar_historial_comisiones(hist)
        return jsonify({"success": True, "historial": hist})

    hist = cargar_historial_comisiones()
    return jsonify({"success": True, "historial": hist})


@app.route("/api/cobranza", methods=["POST"])
def post_cobranza():
    data = request.get_json() or {}
    empresa = data.get("empresa_nombre", "TODAS (Consolidado)")
    vendedor_val = data.get("vendedor_valor", "TODOS")
    vendedor_nom = data.get("vendedor_nombre", "TODOS")
    excluir_inter = data.get("excluir_interempresa", True)

    try:
        res = obtener_cartera_cobranza_engine(
            empresa_nombre=empresa,
            vendedor_valor=vendedor_val,
            vendedor_nombre=vendedor_nom,
            excluir_interempresa=excluir_inter
        )
        return jsonify({"success": True, "data": res})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})


@app.route("/api/cobranza/historial_cliente", methods=["POST"])
def api_cobranza_historial_cliente():
    data = request.get_json() or {}
    cliente_id = data.get("cliente_id")
    cliente_nom = data.get("cliente_nombre", "")
    empresa = data.get("empresa_nombre", "TODAS (Consolidado)")

    try:
        res = obtener_historial_compras_cliente_engine(
            cliente_id=cliente_id,
            cliente_nombre=cliente_nom,
            empresa_nombre=empresa
        )
        return jsonify(res)
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})


@app.route("/api/cobranza/enviar_correo_individual", methods=["POST"])
def api_cobranza_enviar_correo_individual():
    data = request.get_json() or {}
    cliente_id = data.get("cliente_id")
    cliente_nom = data.get("cliente", "Cliente")
    email = data.get("email", "").strip()
    asunto = data.get("asunto", "Recordatorio de Estado de Cuenta - Challenger Automotive")
    cuerpo = data.get("cuerpo", "")
    empresa_nombre = data.get("empresa_nombre", "Challenger Automotive")
    facturas = data.get("facturas", [])

    if not email or "@" not in email:
        return jsonify({"success": False, "error": f"El cliente {cliente_nom} no tiene un correo electrónico válido registrado."})

    try:
        res = enviar_correo_cobranza_odoo(
            cliente_id=cliente_id,
            cliente_nombre=cliente_nom,
            email_destinatario=email,
            asunto=asunto,
            cuerpo_texto=cuerpo,
            empresa_nombre=empresa_nombre,
            facturas=facturas
        )
        return jsonify({
            "success": True,
            "data": res,
            "mensaje": f"¡Correo oficial enviado exitosamente a {cliente_nom} ({email}) registrado en Odoo ERP!"
        })
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})


@app.route("/api/cobranza/enviar_correos_masivos", methods=["POST"])
def api_cobranza_enviar_correos_masivos():
    data = request.get_json() or {}
    destinatarios = data.get("destinatarios", [])
    asunto_base = data.get("asunto", "Recordatorio de Pago y Estado de Cuenta - Challenger Automotive")
    cuerpo_plantilla = data.get("plantilla", "")
    empresa_nombre = data.get("empresa_nombre", "Challenger Automotive")

    if not destinatarios:
        return jsonify({"success": False, "error": "No se seleccionaron clientes destinatarios con correo electrónico."})

    enviados = []
    errores = []

    for d in destinatarios:
        c_id = d.get("cliente_id")
        c_nom = d.get("cliente", "Cliente")
        c_mail = d.get("email", "").strip()
        c_saldo = d.get("saldo_total", 0.0)
        c_dias = d.get("max_dias_retraso", 0)
        c_facs = d.get("facturas_folios", "")
        c_facturas_list = d.get("facturas", [])

        if not c_mail or "@" not in c_mail:
            errores.append({"cliente": c_nom, "error": "Sin correo electrónico válido registrado"})
            continue

        # Personalizar texto
        msg_body = cuerpo_plantilla
        msg_body = msg_body.replace("{cliente}", c_nom)
        msg_body = msg_body.replace("{saldo_total}", f"${float(c_saldo):,.2f}")
        msg_body = msg_body.replace("{dias_vencido}", str(c_dias))
        msg_body = msg_body.replace("{facturas}", str(c_facs))
        msg_body = msg_body.replace("{empresa}", empresa_nombre)
        msg_body = msg_body.replace("{fecha_hoy}", datetime.now().strftime("%d/%m/%Y"))

        try:
            res_envio = enviar_correo_cobranza_odoo(
                cliente_id=c_id,
                cliente_nombre=c_nom,
                email_destinatario=c_mail,
                asunto=asunto_base,
                cuerpo_texto=msg_body,
                empresa_nombre=empresa_nombre,
                facturas=c_facturas_list
            )
            enviados.append({
                "cliente": c_nom,
                "email": c_mail,
                "saldo": c_saldo,
                "dias_retraso": c_dias,
                "message_id": res_envio.get("message_id"),
                "status": "ENVIADO_EXITOSO",
                "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            })
        except Exception as e_ind:
            errores.append({"cliente": c_nom, "email": c_mail, "error": str(e_ind)})

    return jsonify({
        "success": True,
        "total_procesados": len(destinatarios),
        "total_enviados": len(enviados),
        "total_errores": len(errores),
        "enviados": enviados,
        "errores": errores,
        "mensaje": f"¡Campaña de cobranza procesada! Se enviaron {len(enviados)} correos oficiales registrados en Odoo ERP."
    })


@app.route("/api/productos", methods=["POST"])
def post_productos():
    data = request.get_json() or {}
    fecha_ini = data.get("fecha_inicio")
    fecha_fin = data.get("fecha_fin")
    empresa = data.get("empresa_nombre", "TODAS (Consolidado)")
    excluir_inter = data.get("excluir_interempresa", True)

    if not fecha_ini or not fecha_fin:
        anio = datetime.now().year
        mes = datetime.now().month
        import calendar
        _, max_d = calendar.monthrange(anio, mes)
        fecha_ini = f"{anio}-{mes:02d}-01"
        fecha_fin = f"{anio}-{mes:02d}-{max_d:02d}"

    try:
        res = obtener_productos_engine(fecha_ini, fecha_fin, empresa, excluir_inter)
        return jsonify({"success": True, "data": res})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})


@app.route("/api/productos/historico_ml", methods=["POST"])
def post_productos_historico_ml():
    data = request.get_json() or {}
    empresa = data.get("empresa_nombre", "TODAS (Consolidado)")
    excluir_inter = data.get("excluir_interempresa", True)

    try:
        res = obtener_historico_top20_ml(empresa, excluir_interempresa=excluir_inter)
        return jsonify({"success": True, "data": res})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})


@app.route("/api/stock", methods=["POST", "GET"])
def post_stock():
    data = request.get_json(silent=True) or {}
    empresa = data.get("empresa_nombre") or request.args.get("empresa_nombre") or "TODAS (Consolidado)"
    cid = EMPRESAS.get(empresa, {}).get("company_id")
    almacen = data.get("almacen_filtro") or request.args.get("almacen_filtro") or "Todos los Almacenes"
    dias = int(data.get("dias_analisis") or request.args.get("dias_analisis") or request.args.get("dias") or 30)
    excluir_inter = data.get("excluir_interempresa", True)
    if isinstance(excluir_inter, str):
        excluir_inter = excluir_inter.lower() in ("true", "1", "yes")

    try:
        res = obtener_stock_engine(cid, almacen, dias, excluir_inter)
        counts = {}
        for i in res.get("inventario", []):
            c = i.get("clasificacion", "MULTIMARCA")
            counts[c] = counts.get(c, 0) + 1
        print(f"[FLASK API STOCK DEBUG] Clasificaciones: {counts}", flush=True)
        return jsonify({"success": True, "data": res})
    except Exception as e:
        import traceback
        traceback.print_exc()
        demo = generar_datos_demo_stock()
        return jsonify({"success": True, "data": demo, "warning": f"Demostración activada: {e}"})


@app.route("/api/sat/cuentas_por_pagar", methods=["POST", "GET"])
@app.route("/api/sat/facturas", methods=["POST", "GET"])
def post_sat():
    data = request.get_json(silent=True) or {}
    empresa = data.get("empresa_nombre") or request.args.get("empresa_nombre") or "TODAS (Consolidado)"
    anio = data.get("anio") or request.args.get("anio") or "TODOS"
    mes = int(data.get("mes") or request.args.get("mes") or 0)
    tipo = data.get("tipo_comprobante") or request.args.get("tipo_comprobante") or "TODOS"
    direccion = data.get("direccion") or request.args.get("direccion") or "PROVEEDORES"
    try:
        res = procesar_facturas_sat_puras(empresa_nombre=empresa, anio=anio, mes=mes, tipo_comprobante=tipo, direccion=direccion)
        return jsonify({"success": True, "data": res})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)})


@app.route("/api/sat/descargar_xml", methods=["GET"])
def descargar_xml_sat():
    uuid_str = request.args.get("uuid", "").strip()
    if not uuid_str:
        return jsonify({"success": False, "error": "Falta el parámetro uuid."}), 400

    fpath = obtener_xml_por_uuid(uuid_str)
    if not fpath or not os.path.exists(fpath):
        return jsonify({"success": False, "error": f"No se encontró el archivo XML para el UUID {uuid_str}."}), 404

    return send_file(
        fpath,
        as_attachment=True,
        download_name=f"SAT_CFDI_{uuid_str}.xml",
        mimetype="application/xml"
    )


@app.route("/api/sat/descargar_zip", methods=["POST", "GET"])
def descargar_zip_sat():
    data = request.get_json(silent=True) or {}
    uuids = data.get("uuids") or request.args.getlist("uuids")
    empresa = data.get("empresa_nombre") or request.args.get("empresa_nombre") or "TODAS (Consolidado)"
    anio = data.get("anio") or request.args.get("anio") or "TODOS"
    mes = int(data.get("mes") or request.args.get("mes") or 0)
    direccion = data.get("direccion") or request.args.get("direccion") or "PROVEEDORES"

    try:
        import io
        zip_bytes = generar_zip_facturas_sat(uuids=uuids, empresa_nombre=empresa, anio=anio, mes=mes, direccion=direccion)
        filename = f"Facturas_SAT_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
        return send_file(
            io.BytesIO(zip_bytes),
            as_attachment=True,
            download_name=filename,
            mimetype="application/zip"
        )
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/api/sat/detalle_cfdi", methods=["POST", "GET"])
def detalle_cfdi_sat():
    data = request.get_json(silent=True) or {}
    uuid_str = data.get("uuid") or request.args.get("uuid", "").strip()
    if not uuid_str:
        return jsonify({"success": False, "error": "Falta el UUID del comprobante."}), 400

    indice = construir_indice_sat_xmls()
    cfdi = indice.get(uuid_str.upper())
    if not cfdi:
        fpath = obtener_xml_por_uuid(uuid_str)
        if fpath:
            cfdi = parsear_xml_cfdi(fpath)

    if not cfdi:
        return jsonify({"success": False, "error": f"No se encontró el CFDI con UUID {uuid_str}."}), 404

    return jsonify({"success": True, "data": cfdi})


@app.route("/api/pnl", methods=["POST", "GET"])
@app.route("/api/estado_resultados", methods=["POST", "GET"])
def post_estado_resultados():
    data = request.get_json(silent=True) or {}
    empresa = data.get("empresa_nombre", "TODAS (Consolidado)") if request.method == "POST" else request.args.get("empresa_nombre", "TODAS (Consolidado)")
    anio = data.get("anio", 2026) if request.method == "POST" else request.args.get("anio", 2026)
    mes = data.get("mes", 0) if request.method == "POST" else request.args.get("mes", 0)
    try:
        res = procesar_estado_resultados(empresa_nombre=empresa, anio=anio, mes=mes)
        return jsonify({"success": True, "data": res})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)})


@app.route("/api/sat/download", methods=["POST"])
def post_sat_download():
    data = request.get_json() or {}
    cer_p = data.get("cer_path")
    key_p = data.get("key_path")
    pwd   = data.get("password")
    anio  = int(data.get("anio")) if data.get("anio") else None
    mes   = int(data.get("mes")) if data.get("mes") else None

    try:
        res = descargar_cfdis_sat(cer_p, key_p, pwd, anio, mes)
        return jsonify({"success": True, "data": res})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})


@app.route("/api/sat/config", methods=["GET", "POST"])
def get_post_sat_config():
    if request.method == "POST":
        data = request.get_json() or {}
        cfg = guardar_config_fiel(
            data.get("cer_path", ""),
            data.get("key_path", ""),
            data.get("password", "")
        )
        return jsonify({"success": True, "config": cfg})
    else:
        cfg = cargar_config_fiel()
        cfg_safe = {**cfg, "password": "••••••••" if cfg.get("password") else ""}
        return jsonify({"success": True, "config": cfg_safe})


@app.route("/api/ventas_vendedor", methods=["POST"])
def post_ventas_vendedor():
    data = request.get_json() or {}
    empresa = data.get("empresa_nombre", "TODAS (Consolidado)")
    mes = data.get("mes")
    anio = data.get("anio")
    vendedor = data.get("vendedor", "TODOS")
    excluir_inter = data.get("excluir_interempresa", True)
    
    try:
        from ventas_vendedor_engine import obtener_ventas_por_vendedor_engine
        res = obtener_ventas_por_vendedor_engine(
            empresa_nombre=empresa,
            mes=mes,
            anio=anio,
            vendedor_filtro=vendedor,
            excluir_interempresa=excluir_inter
        )
        return jsonify({"success": True, "data": res})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)})


@app.route("/api/ventas_vendedor_3meses", methods=["POST"])
def post_ventas_vendedor_3meses():
    data = request.get_json() or {}
    empresa = data.get("empresa_nombre", "TODAS (Consolidado)")
    mes = data.get("mes")
    anio = data.get("anio")
    excluir_inter = data.get("excluir_interempresa", True)
    
    try:
        from ventas_vendedor_engine import obtener_reporte_3meses_vendedores_engine
        res = obtener_reporte_3meses_vendedores_engine(
            empresa_nombre=empresa,
            mes_fin=mes,
            anio_fin=anio,
            excluir_interempresa=excluir_inter
        )
        return jsonify({"success": True, "data": res})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)})


@app.route("/api/cobranza/historial_cliente", methods=["POST"])
def post_cobranza_historial_cliente():
    data = request.get_json() or {}
    cliente_id = data.get("cliente_id")
    cliente_nombre = data.get("cliente_nombre", "")
    empresa = data.get("empresa_nombre", "TODAS (Consolidado)")
    
    try:
        from odoo_sync_engine import obtener_historial_compras_cliente_engine
        res = obtener_historial_compras_cliente_engine(
            cliente_id=cliente_id,
            cliente_nombre=cliente_nombre,
            empresa_nombre=empresa
        )
        return jsonify(res)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)})


@app.route("/api/export_excel", methods=["POST"])
def export_excel():
    data = request.get_json() or {}
    tipo = data.get("tipo", "stock")
    items = data.get("items", [])

    try:
        import openpyxl
        from excel_styler import (
            aplicar_estilo_banner, agregar_seccion_kpis, agregar_tabla_ejecutiva,
            COLOR_SLATE_DARK, COLOR_PURPLE_HEADER, COLOR_BLUE_HEADER, COLOR_AMBER_HEADER, COLOR_EMERALD_HEADER
        )
    except ImportError as e:
        return jsonify({"success": False, "error": f"Error importando estilos: {str(e)}"})

    wb = openpyxl.Workbook()
    ws = wb.active

    if tipo == "productos":
        # ── HOJA 1: RANKING DE PRODUCTOS ──────────────────────────────────
        ws.title = "Ranking Productos"
        tot_monto = sum(float(i.get("subtotal", 0) or 0) for i in items)
        tot_cant = sum(float(i.get("cantidad", 0) or 0) for i in items)
        
        aplicar_estilo_banner(
            ws,
            titulo="CHALLENGER AUTOMOTIVE — TOP PRODUCTOS MÁS VENDIDOS",
            subtitulo="Ranking consolidado de ventas, volumen y participación de mercado",
            num_columnas=9,
            tema_color=COLOR_SLATE_DARK
        )
        
        kpis_rank = [
            {"label": "Total Facturado", "value": f"${tot_monto:,.2f}", "text_color": "065F46"},
            {"label": "Unidades Vendidas", "value": f"{int(tot_cant):,} uds", "text_color": "1E40AF"},
            {"label": "Variedad Códigos", "value": f"{len(items)} productos", "text_color": "581C87"}
        ]
        next_r = agregar_seccion_kpis(ws, kpis_rank, start_row=4)
        
        headers_rank = ["Posición", "Producto / Código", "Categoría / Marca", "Cantidad Vendida", "Unidad", "Total Facturado ($)", "Precio Promedio ($)", "Facturas #", "% del Total"]
        rows_rank = [
            [
                i.get("posicion"), i.get("nombre"), i.get("clasificacion"),
                i.get("cantidad"), i.get("uom"), i.get("subtotal"),
                i.get("precio_promedio"), i.get("num_facturas"), i.get("pct_monto")
            ]
            for i in items
        ]
        formatos_rank = {1: "center", 2: "left", 3: "badge", 4: "number", 5: "center", 6: "currency", 7: "currency", 8: "center", 9: "percent"}
        totales_rank = ["TOTALES", f"{len(items)} Códigos", "—", tot_cant, "—", tot_monto, "—", "—", "100.0%"]
        
        agregar_tabla_ejecutiva(ws, next_r, headers_rank, rows_rank, formatos_rank, tema_color=COLOR_SLATE_DARK, totales_row=totales_rank)

        # ── HOJA 2 & 3: MACHINE LEARNING & HISTÓRICO SI SE PROPORCIONÓ ───
        ml_data = data.get("ml_data")
        if ml_data:
            # HOJA 2: MACHINE LEARNING FORECAST
            mes_pron = ml_data.get("mes_pronosticado", "Próximo Mes")
            ws_ml = wb.create_sheet(title=f"Pronóstico {mes_pron}")
            
            aplicar_estilo_banner(
                ws_ml,
                titulo=f"CHALLENGER AI — PRONÓSTICO DE MACHINE LEARNING ({mes_pron.upper()})",
                subtitulo=f"Modelo predictivo XGBoost Regressor (Gradient Boosting) | Demanda futura y buffer de stock",
                num_columnas=11,
                tema_color=COLOR_PURPLE_HEADER
            )
            
            preds = ml_data.get("predicciones_ml", [])
            tot_pred_u = sum(float(p.get("pred_unidades", 0) or 0) for p in preds)
            tot_pred_ing = sum(float(p.get("pred_ingreso", 0) or 0) for p in preds)
            tot_buf = sum(float(p.get("stock_sugerido", 0) or 0) for p in preds)
            conf_avg = (sum(float(p.get("confianza_score", 0) or 0) for p in preds) / len(preds)) if preds else 0.0
            
            kpis_ml = [
                {"label": f"Demanda Proyectada ({mes_pron})", "value": f"{int(tot_pred_u):,} uds", "text_color": "581C87"},
                {"label": f"Ingreso Estimado ({mes_pron})", "value": f"${tot_pred_ing:,.2f}", "text_color": "065F46"},
                {"label": "Stock Buffer Sugerido (+25%)", "value": f"{int(tot_buf):,} uds", "text_color": "1E40AF"},
                {"label": "Confianza Modelo IA", "value": f"{conf_avg:.1f}%", "text_color": "92400E"}
            ]
            next_r_ml = agregar_seccion_kpis(ws_ml, kpis_ml, start_row=4)
            
            headers_ml = ["Pos", "Producto / Código", "Marca", "Mes Anterior (Uds)", "Promedio Histórico", f"Pronóstico {mes_pron} (Uds)", f"Ingreso Estimado ($)", "Buffer Almacén", "Tendencia IA", "Confianza", "Estrategia Comercial IA"]
            rows_ml = [
                [
                    p.get("ranking"), p.get("producto"), p.get("categoria"),
                    p.get("ult_mes_unidades"), p.get("promedio_mensual"),
                    p.get("pred_unidades"), p.get("pred_ingreso"),
                    p.get("stock_sugerido"), p.get("tendencia_flag"),
                    p.get("confianza_score"), p.get("explicacion_ia")
                ]
                for p in preds
            ]
            formatos_ml = {
                1: "center", 2: "left", 3: "badge", 4: "number", 5: "number",
                6: "number", 7: "currency", 8: "number", 9: "badge", 10: "percent", 11: "left"
            }
            totales_ml = ["TOTALES", f"{len(preds)} Productos", "—", "—", "—", tot_pred_u, tot_pred_ing, tot_buf, "—", f"{conf_avg:.1f}%", "Proyección Consolidada"]
            
            agregar_tabla_ejecutiva(ws_ml, next_r_ml, headers_ml, rows_ml, formatos_ml, tema_color=COLOR_PURPLE_HEADER, totales_row=totales_ml)

            # HOJA 3: MATRIZ HISTÓRICA MULTI-MES
            periodos = ml_data.get("periodos_historico", [])
            if periodos:
                ws_hist = wb.create_sheet(title="Histórico Multi-mes")
                aplicar_estilo_banner(
                    ws_hist,
                    titulo="CHALLENGER AUTOMOTIVE — MATRIZ HISTÓRICA DE VENTAS (CATÁLOGO COMPLETO)",
                    subtitulo=f"Comportamiento de facturación mensual en los últimos {len(periodos)} meses para todos los productos",
                    num_columnas=len(periodos) + 5,
                    tema_color=COLOR_BLUE_HEADER
                )
                
                headers_hist = ["Posición", "Producto / Código", "Marca"] + [f"{per} ($)" for per in periodos] + ["Total Histórico ($)", "Crecimiento MoM (%)"]
                rows_hist = []
                prods_hist = ml_data.get("matriz_historica") or ml_data.get("top20_historico", [])
                for p in prods_hist:
                    r_vals = [p.get("ranking"), p.get("producto"), p.get("categoria")]
                    for per in periodos:
                        sub_per = p.get("ventas_por_mes", {}).get(per, {}).get("subtotal", 0.0)
                        r_vals.append(sub_per)
                    r_vals.append(p.get("total_subtotal"))
                    r_vals.append(p.get("mom_growth"))
                    rows_hist.append(r_vals)
                
                formatos_hist = {1: "center", 2: "left", 3: "badge"}
                for c_i in range(4, len(periodos) + 5):
                    formatos_hist[c_i] = "currency"
                formatos_hist[len(periodos) + 5] = "percent"
                
                agregar_tabla_ejecutiva(ws_hist, 4, headers_hist, rows_hist, formatos_hist, tema_color=COLOR_BLUE_HEADER)

            # HOJA 4: RESUMEN DE AJUSTES Y DESCUENTOS
            met_aj = ml_data.get("metricas_ajuste", {})
            ws_aj = wb.create_sheet(title="Ajustes y Descuentos")
            aplicar_estilo_banner(
                ws_aj,
                titulo="CHALLENGER AUTOMOTIVE — ANÁLISIS DE AJUSTES Y DESCUENTOS",
                subtitulo="Auditoría comercial de descuentos otorgados vs subtotal facturado bruto",
                num_columnas=4,
                tema_color=COLOR_AMBER_HEADER
            )
            headers_aj = ["Concepto Financiero", "Importe ($ MXN)", "Tasa (%)", "Descripción / Impacto"]
            rows_aj = [
                ["Total Descuentos Concedidos", met_aj.get("total_descuentos_otorgados", 0), f"{met_aj.get('pct_descuento_promedio', 0):.2f}%", "Descuentos comerciales directos en factura"],
                ["Subtotal Facturado Bruto", met_aj.get("total_subtotal_bruto", 0), "100.0%", "Valoración total a precio de lista base"],
                ["Total Facturado Neto Efectivo", met_aj.get("total_facturado_neto", 0), f"{100.0 - met_aj.get('pct_descuento_promedio', 0):.2f}%", "Ingreso neto efectivo cobrado en Odoo"]
            ]
            formatos_aj = {1: "left", 2: "currency", 3: "center", 4: "left"}
            agregar_tabla_ejecutiva(ws_aj, 4, headers_aj, rows_aj, formatos_aj, tema_color=COLOR_AMBER_HEADER)

    elif tipo == "productos_ml":
        # ── EXPORTACIÓN DEDICADA DE PRONÓSTICO MACHINE LEARNING ──────────
        ml_data = data.get("ml_data") or {}
        mes_pron = ml_data.get("mes_pronosticado", "Próximo Mes")
        ws.title = f"Pronóstico ML {mes_pron}"
        
        aplicar_estilo_banner(
            ws,
            titulo=f"CHALLENGER AI — PRONÓSTICO DE MACHINE LEARNING ({mes_pron.upper()})",
            subtitulo=f"Modelo predictivo XGBoost Regressor (Gradient Boosting) | Demanda futura y buffer de stock",
            num_columnas=11,
            tema_color=COLOR_PURPLE_HEADER
        )
        
        preds = ml_data.get("predicciones_ml") or items or []
        tot_pred_u = sum(float(p.get("pred_unidades", 0) or 0) for p in preds)
        tot_pred_ing = sum(float(p.get("pred_ingreso", 0) or 0) for p in preds)
        tot_buf = sum(float(p.get("stock_sugerido", 0) or 0) for p in preds)
        conf_avg = (sum(float(p.get("confianza_score", 0) or 0) for p in preds) / len(preds)) if preds else 0.0
        
        kpis_ml = [
            {"label": f"Demanda Proyectada ({mes_pron})", "value": f"{int(tot_pred_u):,} uds", "text_color": "581C87"},
            {"label": f"Ingreso Estimado ({mes_pron})", "value": f"${tot_pred_ing:,.2f}", "text_color": "065F46"},
            {"label": "Stock Buffer Sugerido (+25%)", "value": f"{int(tot_buf):,} uds", "text_color": "1E40AF"},
            {"label": "Confianza Modelo IA", "value": f"{conf_avg:.1f}%", "text_color": "92400E"}
        ]
        next_r_ml = agregar_seccion_kpis(ws, kpis_ml, start_row=4)
        
        headers_ml = ["Pos", "Producto / Código", "Marca", "Mes Anterior (Uds)", "Promedio Histórico", f"Pronóstico {mes_pron} (Uds)", f"Ingreso Estimado ($)", "Buffer Almacén", "Tendencia IA", "Confianza", "Estrategia Comercial IA"]
        rows_ml = [
            [
                p.get("ranking"), p.get("producto"), p.get("categoria"),
                p.get("ult_mes_unidades"), p.get("promedio_mensual"),
                p.get("pred_unidades"), p.get("pred_ingreso"),
                p.get("stock_sugerido"), p.get("tendencia_flag"),
                p.get("confianza_score"), p.get("explicacion_ia")
            ]
            for p in preds
        ]
        formatos_ml = {
            1: "center", 2: "left", 3: "badge", 4: "number", 5: "number",
            6: "number", 7: "currency", 8: "number", 9: "badge", 10: "percent", 11: "left"
        }
        totales_ml = ["TOTALES", f"{len(preds)} Productos", "—", "—", "—", tot_pred_u, tot_pred_ing, tot_buf, "—", f"{conf_avg:.1f}%", "Proyección Consolidada"]
        
        agregar_tabla_ejecutiva(ws, next_r_ml, headers_ml, rows_ml, formatos_ml, tema_color=COLOR_PURPLE_HEADER, totales_row=totales_ml)

    elif tipo == "productos_historico":
        # ── EXPORTACIÓN DEDICADA DE MATRIZ HISTÓRICA MULTI-MES ────────────
        ml_data = data.get("ml_data") or {}
        periodos = ml_data.get("periodos_historico", [])
        prods_hist = ml_data.get("matriz_historica") or ml_data.get("top20_historico") or items or []
        ws.title = "Matriz Histórica Ventas"
        
        aplicar_estilo_banner(
            ws,
            titulo="CHALLENGER AUTOMOTIVE — MATRIZ HISTÓRICA DE VENTAS (CATÁLOGO COMPLETO)",
            subtitulo=f"Comportamiento de facturación mensual en los últimos {len(periodos)} meses para todos los productos",
            num_columnas=len(periodos) + 5,
            tema_color=COLOR_BLUE_HEADER
        )
        
        headers_hist = ["Posición", "Producto / Código", "Marca"] + [f"{per} ($)" for per in periodos] + ["Total Histórico ($)", "Crecimiento MoM (%)"]
        rows_hist = []
        for p in prods_hist:
            r_vals = [p.get("ranking"), p.get("producto"), p.get("categoria")]
            for per in periodos:
                sub_per = p.get("ventas_por_mes", {}).get(per, {}).get("subtotal", 0.0)
                r_vals.append(sub_per)
            r_vals.append(p.get("total_subtotal"))
            r_vals.append(p.get("mom_growth"))
            rows_hist.append(r_vals)
        
        formatos_hist = {1: "center", 2: "left", 3: "badge"}
        for c_i in range(4, len(periodos) + 5):
            formatos_hist[c_i] = "currency"
        formatos_hist[len(periodos) + 5] = "percent"
        
        agregar_tabla_ejecutiva(ws, 4, headers_hist, rows_hist, formatos_hist, tema_color=COLOR_BLUE_HEADER)

    elif tipo == "cobranza":
        # ── REPORTE DE COBRANZA & CARTERA ─────────────────────────────────
        ws.title = "Facturas por Cobrar"
        tot_deuda = sum(float(i.get("saldo_pendiente", 0) or 0) for i in items)
        tot_al_dia = sum(float(i.get("saldo_pendiente", 0) or 0) for i in items if (i.get("dias_retraso", 0) or 0) <= 0)
        tot_reciente = sum(float(i.get("saldo_pendiente", 0) or 0) for i in items if 0 < (i.get("dias_retraso", 0) or 0) <= 60)
        tot_pasada = sum(float(i.get("saldo_pendiente", 0) or 0) for i in items if 60 < (i.get("dias_retraso", 0) or 0) <= 100)
        tot_atrasada = sum(float(i.get("saldo_pendiente", 0) or 0) for i in items if (i.get("dias_retraso", 0) or 0) > 100)
        
        aplicar_estilo_banner(
            ws,
            titulo="CHALLENGER AUTOMOTIVE — CARTERA DE COBRANZA Y GESTIÓN DE COBRO",
            subtitulo="Antigüedad de saldos (0-60d, 60-100d, >100d), contacto de clientes y asignación de vendedor",
            num_columnas=12,
            tema_color=COLOR_EMERALD_HEADER
        )
        
        kpis_cob = [
            {"label": "Saldo Total Cartera", "value": f"${tot_deuda:,.2f}", "text_color": "991B1B"},
            {"label": "🟢 Reciente (0-60 días)", "value": f"${tot_reciente:,.2f}", "text_color": "065F46"},
            {"label": "🟡 Pasada (60-100 días)", "value": f"${tot_pasada:,.2f}", "text_color": "92400E"},
            {"label": "🚨 Atrasada (+100 días)", "value": f"${tot_atrasada:,.2f}", "text_color": "991B1B"}
        ]
        next_r = agregar_seccion_kpis(ws, kpis_cob, start_row=4)
        
        headers_cob = ["Folio Factura", "Cliente / Razón Social", "Teléfono", "Email", "Vendedor", "Fecha Fac", "Vencimiento", "Días Retraso", "Rango Cobranza", "Total Factura ($)", "Abonado ($)", "Saldo Pendiente ($)"]
        rows_cob = []
        for i in items:
            d = i.get("dias_retraso", 0) or 0
            if d > 100: r_txt = "🚨 Atrasada (>100d)"
            elif d > 60: r_txt = "🟡 Pasada (60-100d)"
            elif d > 0: r_txt = "🟢 Reciente (0-60d)"
            else: r_txt = "✓ Al Día"
            
            rows_cob.append([
                i.get("folio"), i.get("cliente"), i.get("telefono") or "—", i.get("email") or "—",
                i.get("vendedor"), i.get("fecha_factura"), i.get("fecha_vencimiento"), d,
                r_txt, i.get("total_factura"), i.get("monto_abonado"), i.get("saldo_pendiente")
            ])
            
        formatos_cob = {1: "center", 2: "left", 3: "center", 4: "left", 5: "left", 6: "center", 7: "center", 8: "center", 9: "badge", 10: "currency", 11: "currency", 12: "currency"}
        tot_fac = sum(float(i.get("total_factura", 0) or 0) for i in items)
        tot_ab = sum(float(i.get("monto_abonado", 0) or 0) for i in items)
        totales_cob = ["TOTALES", f"{len(items)} Facturas", "—", "—", "—", "—", "—", "—", "—", tot_fac, tot_ab, tot_deuda]
        
        agregar_tabla_ejecutiva(ws, next_r, headers_cob, rows_cob, formatos_cob, tema_color=COLOR_EMERALD_HEADER, totales_row=totales_cob)

        # ── HOJA 2: CICLOS VENTA VS RECUPERACIÓN AL SIGUIENTE MES ──────────
        ciclos_data = data.get("ciclos_data") or {}
        meses_ciclos = ciclos_data.get("meses") or []
        if meses_ciclos:
            ws_ciclos = wb.create_sheet(title="Ciclos Venta y Recuperación")
            aplicar_estilo_banner(
                ws_ciclos,
                titulo="CHALLENGER AUTOMOTIVE — CICLOS DE VENTA Y RECUPERACIÓN AL SIGUIENTE MES",
                subtitulo="Relación entre importe vendido en el mes, compromiso a recuperar al siguiente mes y tasa de efectividad de cobranza",
                num_columnas=9,
                tema_color=COLOR_BLUE_HEADER
            )
            
            headers_ciclos = ["Mes de Venta", "Mes a Recuperar (Siguiente)", "Total Vendido ($)", "A Recuperar Mes Sig. ($)", "Ya Cobrado ($)", "Saldo Pendiente ($)", "Tasa Recuperación (%)", "Facturas #", "Clientes #"]
            rows_ciclos = [
                [
                    m.get("mes_venta"), m.get("mes_siguiente"), m.get("total_vendido"),
                    m.get("a_recuperar_mes_sig"), m.get("ya_recuperado"), m.get("saldo_pendiente"),
                    (m.get("tasa_recuperacion", 0) or 0) / 100.0, m.get("facturas_count"), m.get("clientes_count")
                ]
                for m in meses_ciclos
            ]
            formatos_ciclos = {1: "center", 2: "center", 3: "currency", 4: "currency", 5: "currency", 6: "currency", 7: "percent", 8: "number", 9: "number"}
            tot_v = sum(float(m.get("total_vendido", 0) or 0) for m in meses_ciclos)
            tot_r = sum(float(m.get("a_recuperar_mes_sig", 0) or 0) for m in meses_ciclos)
            tot_c = sum(float(m.get("ya_recuperado", 0) or 0) for m in meses_ciclos)
            tot_p = sum(float(m.get("saldo_pendiente", 0) or 0) for m in meses_ciclos)
            tasa_glob = (tot_c / tot_v) if tot_v > 0 else 0.0
            totales_ciclos = ["TOTALES", f"{len(meses_ciclos)} Meses", tot_v, tot_r, tot_c, tot_p, tasa_glob, sum(m.get("facturas_count", 0) for m in meses_ciclos), "—"]
            
            agregar_tabla_ejecutiva(ws_ciclos, 4, headers_ciclos, rows_ciclos, formatos_ciclos, tema_color=COLOR_BLUE_HEADER, totales_row=totales_ciclos)

    elif tipo == "cobrado_mes":
        # ── REPORTE DE COBRADO & RECUPERACIÓN MENSUAL ─────────────────────
        cob_payload = data.get("cobrado_mes_data") or {}
        meses_data = cob_payload.get("meses") or []
        res_glob = cob_payload.get("resumen_global") or {}
        per_filtro = data.get("periodo_filtro", "TODOS")
        vend_filtro = data.get("vendedor_filtro", "TODOS")

        # HOJA 1: RESUMEN HISTÓRICO DE COBRANZA MENSUAL
        ws.title = "Matriz Mensual Histórico"
        aplicar_estilo_banner(
            ws,
            titulo="CHALLENGER AUTOMOTIVE — COBRADO & RECUPERACIÓN POR MES",
            subtitulo=f"Filtro: {per_filtro} | Vendedor: {vend_filtro} | Auditoría integral de vencimientos y pagos en Odoo ERP",
            num_columnas=8,
            tema_color=COLOR_EMERALD_HEADER
        )

        tot_est_hist = sum(float(m.get("cobranza_estimada_vence", 0) or 0) for m in meses_data)
        tot_rec_hist = sum(float(m.get("cobrado_de_vencidas", 0) or 0) for m in meses_data)
        tot_pen_hist = sum(float(m.get("saldo_pendiente_vencidas", 0) or 0) for m in meses_data)
        tot_ing_hist = sum(float(m.get("cobrado_real_pagos", m.get("total_cobrado", 0)) or 0) for m in meses_data)
        tot_fac_hist = sum(int(m.get("facturas_vencen_count", 0) or 0) for m in meses_data)
        tot_pag_hist = sum(int(m.get("num_pagos_recibidos", m.get("num_pagos", 0)) or 0) for m in meses_data)
        pct_ef_hist = (tot_rec_hist / tot_est_hist * 100.0) if tot_est_hist > 0 else 100.0

        kpis_cm = [
            {"label": "Total Estimado a Cobrar", "value": f"${tot_est_hist:,.2f}", "text_color": "1E40AF"},
            {"label": "Recuperación de Facturas", "value": f"${tot_rec_hist:,.2f}", "text_color": "065F46"},
            {"label": "Saldo Insoluto Pendiente", "value": f"${tot_pen_hist:,.2f}", "text_color": "991B1B"},
            {"label": "Ingresos a Bancos Global", "value": f"${tot_ing_hist:,.2f}", "text_color": "581C87"}
        ]
        next_r_cm = agregar_seccion_kpis(ws, kpis_cm, start_row=4)

        headers_cm = ["Mes de Cobro / Período", "Estimado a Cobrar ($)", "Facturas (#)", "Recuperado del Mes ($)", "Pendiente ($)", "% Cumplimiento", "Ingresos Bancos ($)", "Pagos (#)"]
        rows_cm = [
            [
                m.get("mes_nombre"), m.get("cobranza_estimada_vence", 0.0), m.get("facturas_vencen_count", 0),
                m.get("cobrado_de_vencidas", 0.0), m.get("saldo_pendiente_vencidas", 0.0),
                (m.get("pct_cumplimiento_vence", 0) or 0) / 100.0,
                m.get("cobrado_real_pagos", m.get("total_cobrado", 0.0)),
                m.get("num_pagos_recibidos", m.get("num_pagos", 0))
            ]
            for m in meses_data
        ]
        formatos_cm = {1: "center", 2: "currency", 3: "number", 4: "currency", 5: "currency", 6: "percent", 7: "currency", 8: "number"}
        totales_cm = [
            "TOTALES", tot_est_hist, tot_fac_hist, tot_rec_hist, tot_pen_hist,
            pct_ef_hist / 100.0, tot_ing_hist, tot_pag_hist
        ]
        agregar_tabla_ejecutiva(ws, next_r_cm, headers_cm, rows_cm, formatos_cm, tema_color=COLOR_EMERALD_HEADER, totales_row=totales_cm)

        # HOJA 2: CLIENTES CON VENCIMIENTO EN EL MES
        ws_cli = wb.create_sheet(title="Clientes con Vencimiento")
        aplicar_estilo_banner(
            ws_cli,
            titulo="CHALLENGER AUTOMOTIVE — CLIENTES CON VENCIMIENTO EN EL PERÍODO",
            subtitulo=f"Período: {per_filtro} | Vendedor: {vend_filtro} | Facturación vencida, abonos y saldo insoluto",
            num_columnas=10,
            tema_color=COLOR_PURPLE_HEADER
        )

        clientes_export = []
        if per_filtro == "TODOS":
            cmap = {}
            for m in meses_data:
                for c in m.get("clientes", []):
                    cid = c.get("cliente_id") or c.get("cliente_nombre")
                    if cid not in cmap:
                        cmap[cid] = {
                            "cliente_nombre": c.get("cliente_nombre"),
                            "rfc": c.get("rfc"),
                            "vendedor": c.get("vendedor"),
                            "total_estimado_vence": 0.0,
                            "total_cobrado_vencidas": 0.0,
                            "saldo_pendiente_vencidas": 0.0,
                            "num_facturas_vence": 0,
                            "mes": "Histórico Consolidado"
                        }
                    cmap[cid]["total_estimado_vence"] += float(c.get("total_estimado_vence", 0) or 0)
                    cmap[cid]["total_cobrado_vencidas"] += float(c.get("total_cobrado_vencidas", 0) or 0)
                    cmap[cid]["saldo_pendiente_vencidas"] += float(c.get("saldo_pendiente_vencidas", 0) or 0)
                    cmap[cid]["num_facturas_vence"] += int(c.get("num_facturas_vence", 0) or 0)
            clientes_export = [c for c in cmap.values() if c["total_estimado_vence"] > 0.001 or c["num_facturas_vence"] > 0]
        else:
            m_target = next((m for m in meses_data if m.get("mes_key") == per_filtro), None)
            if m_target:
                for c in m_target.get("clientes", []):
                    if float(c.get("total_estimado_vence", 0) or 0) > 0.001 or int(c.get("num_facturas_vence", 0) or 0) > 0:
                        c_copy = dict(c)
                        c_copy["mes"] = m_target.get("mes_nombre")
                        clientes_export.append(c_copy)

        if vend_filtro and vend_filtro != "TODOS":
            clientes_export = [c for c in clientes_export if normalizar_texto(c.get("vendedor", "")) == normalizar_texto(vend_filtro)]

        headers_cl = ["#", "Cliente / Razón Social", "RFC", "Asesor Comercial", "Mes / Período", "🎯 Vence en Mes ($)", "🟢 Abonado ($)", "⏳ Pendiente ($)", "% Cumplimiento", "Facturas #"]
        rows_cl = []
        for idx, c in enumerate(sorted(clientes_export, key=lambda x: float(x.get("total_estimado_vence", 0) or 0), reverse=True)):
            est_c = float(c.get("total_estimado_vence", 0) or 0)
            cob_c = float(c.get("total_cobrado_vencidas", 0) or 0)
            pen_c = float(c.get("saldo_pendiente_vencidas", 0) or 0)
            pct_c = (cob_c / est_c) if est_c > 0 else 0.0
            rows_cl.append([
                idx + 1, c.get("cliente_nombre"), c.get("rfc") or "—",
                c.get("vendedor") or "Sin Asignar", c.get("mes"),
                est_c, cob_c, pen_c, pct_c, c.get("num_facturas_vence", 0)
            ])
        formatos_cl = {1: "center", 2: "left", 3: "center", 4: "left", 5: "center", 6: "currency", 7: "currency", 8: "currency", 9: "percent", 10: "number"}
        tot_cl_est = sum(float(c.get("total_estimado_vence", 0) or 0) for c in clientes_export)
        tot_cl_cob = sum(float(c.get("total_cobrado_vencidas", 0) or 0) for c in clientes_export)
        tot_cl_sal = sum(float(c.get("saldo_pendiente_vencidas", 0) or 0) for c in clientes_export)
        tot_cl_fac = sum(int(c.get("num_facturas_vence", 0) or 0) for c in clientes_export)
        pct_cl_tot = (tot_cl_cob / tot_cl_est) if tot_cl_est > 0 else 0.0
        totales_cl = ["TOTALES", f"{len(clientes_export)} Clientes", "—", "—", "—", tot_cl_est, tot_cl_cob, tot_cl_sal, pct_cl_tot, tot_cl_fac]
        agregar_tabla_ejecutiva(ws_cli, 4, headers_cl, rows_cl, formatos_cl, tema_color=COLOR_PURPLE_HEADER, totales_row=totales_cl)

        # HOJA 3: FACTURAS QUE VENCEN EN EL PERÍODO
        ws_facs = wb.create_sheet(title="Facturas que Vencen")
        aplicar_estilo_banner(
            ws_facs,
            titulo="CHALLENGER AUTOMOTIVE — FACTURAS CON FECHA DE VENCIMIENTO EN EL PERÍODO",
            subtitulo=f"Período: {per_filtro} | Vendedor: {vend_filtro} | Relación individual de comprobantes",
            num_columnas=11,
            tema_color=COLOR_BLUE_HEADER
        )

        facturas_export = []
        for m in (meses_data if per_filtro == "TODOS" else [m for m in meses_data if m.get("mes_key") == per_filtro]):
            for f in m.get("facturas_vencen", []):
                if vend_filtro == "TODOS" or normalizar_texto(f.get("vendedor", "")) == normalizar_texto(vend_filtro):
                    f_c = dict(f)
                    f_c["mes_vencimiento"] = m.get("mes_nombre")
                    facturas_export.append(f_c)

        headers_fc = ["#", "Folio", "Cliente / Razón Social", "RFC", "Asesor Comercial", "Emisión", "Vencimiento", "Total Factura ($)", "Monto Abonado ($)", "Saldo Pendiente ($)", "Estatus"]
        rows_fc = []
        for idx, f in enumerate(sorted(facturas_export, key=lambda x: (x.get("fecha_vencimiento", ""), x.get("saldo_pendiente", 0)), reverse=True)):
            tot_f = float(f.get("total_factura", 0) or 0)
            sal_f = float(f.get("saldo_pendiente", 0) or 0)
            abo_f = float(f.get("monto_abonado", tot_f - sal_f) or 0)
            st_f = "🟢 Pagada" if sal_f <= 0.01 else ("🟡 Parcial" if abo_f > 0 else "🚨 Pendiente")
            rows_fc.append([
                idx + 1, f.get("folio"), f.get("cliente"), f.get("rfc") or "—",
                f.get("vendedor") or "Sin Asignar", f.get("fecha_emision"), f.get("fecha_vencimiento"),
                tot_f, abo_f, sal_f, st_f
            ])
        formatos_fc = {1: "center", 2: "center", 3: "left", 4: "center", 5: "left", 6: "center", 7: "center", 8: "currency", 9: "currency", 10: "currency", 11: "badge"}
        tot_fc_tot = sum(float(f.get("total_factura", 0) or 0) for f in facturas_export)
        tot_fc_abo = sum(float(f.get("monto_abonado", 0) or 0) for f in facturas_export)
        tot_fc_sal = sum(float(f.get("saldo_pendiente", 0) or 0) for f in facturas_export)
        totales_fc = ["TOTALES", f"{len(facturas_export)} Facturas", "—", "—", "—", "—", "—", tot_fc_tot, tot_fc_abo, tot_fc_sal, "—"]
        agregar_tabla_ejecutiva(ws_facs, 4, headers_fc, rows_fc, formatos_fc, tema_color=COLOR_BLUE_HEADER, totales_row=totales_fc)

        # HOJA 4: RELACIÓN DE PAGOS BANCARIOS
        ws_pag = wb.create_sheet(title="Pagos Bancarios Recibidos")
        aplicar_estilo_banner(
            ws_pag,
            titulo="CHALLENGER AUTOMOTIVE — AUDITORÍA DE PAGOS BANCARIOS RECIBIDOS",
            subtitulo=f"Período: {per_filtro} | Vendedor: {vend_filtro} | Folios, diarios contables y conceptos en Odoo ERP",
            num_columnas=9,
            tema_color=COLOR_SLATE_DARK
        )

        pagos_export = []
        for m in (meses_data if per_filtro == "TODOS" else [m for m in meses_data if m.get("mes_key") == per_filtro]):
            for p in m.get("pagos_recibidos", []):
                if vend_filtro == "TODOS" or normalizar_texto(p.get("vendedor", "")) == normalizar_texto(vend_filtro):
                    pagos_export.append({
                        "mes": m.get("mes_nombre"),
                        "fecha": p.get("fecha"),
                        "folio": p.get("folio"),
                        "cliente": p.get("cliente"),
                        "rfc": p.get("rfc"),
                        "vendedor": p.get("vendedor"),
                        "diario": p.get("diario"),
                        "memo": p.get("memo"),
                        "monto": p.get("monto")
                    })

        headers_pg = ["#", "Fecha Pago", "Folio Pago", "Cliente", "RFC", "Vendedor", "Diario / Banco", "Concepto / Factura", "Monto Cobrado ($)"]
        rows_pg = [
            [
                idx + 1, p.get("fecha"), p.get("folio"), p.get("cliente"),
                p.get("rfc") or "—", p.get("vendedor") or "—", p.get("diario"),
                p.get("memo") or "—", p.get("monto")
            ]
            for idx, p in enumerate(sorted(pagos_export, key=lambda x: x.get("fecha", ""), reverse=True))
        ]
        formatos_pg = {1: "center", 2: "center", 3: "center", 4: "left", 5: "center", 6: "left", 7: "left", 8: "left", 9: "currency"}
        tot_pg_monto = sum(float(p.get("monto", 0) or 0) for p in pagos_export)
        totales_pg = ["TOTALES", f"{len(pagos_export)} Pagos", "—", "—", "—", "—", "—", "—", tot_pg_monto]
        agregar_tabla_ejecutiva(ws_pag, 4, headers_pg, rows_pg, formatos_pg, tema_color=COLOR_SLATE_DARK, totales_row=totales_pg)

    elif tipo in ["historial_cliente", "compras_cliente"]:
        # ── REPORTE DE HISTORIAL DE COMPRAS & FACTURACIÓN DEL CLIENTE ─────
        hist_data = data.get("hist_data") or {}
        cliente_nom = hist_data.get("cliente_nombre", "Cliente")
        rfc = hist_data.get("rfc", "SIN RFC")
        vendedor = hist_data.get("vendedor", "Sin Asignar")
        kpis = hist_data.get("kpis", {})
        facturas = hist_data.get("facturas", items or [])

        # HOJA 1: RESUMEN DE FACTURAS Y COBRANZA
        ws.title = "Historial Facturas"
        aplicar_estilo_banner(
            ws,
            titulo=f"ESTADO DE CUENTA Y COMPRAS — {cliente_nom.upper()}",
            subtitulo=f"RFC: {rfc} | Vendedor Asignado: {vendedor} | Auditoría integral de ventas y pagos en Odoo ERP",
            num_columnas=11,
            tema_color=COLOR_EMERALD_HEADER
        )

        tot_comprado = float(kpis.get("total_comprado", sum(float(f.get("total", 0) or 0) for f in facturas)) or 0)
        tot_pagado = float(kpis.get("total_pagado", sum(float(f.get("monto_pagado", 0) or 0) for f in facturas)) or 0)
        tot_saldo = float(kpis.get("total_saldo_pendiente", sum(float(f.get("saldo_pendiente", 0) or 0) for f in facturas)) or 0)
        pct_cub = (tot_pagado / tot_comprado * 100) if tot_comprado > 0 else 100.0

        kpis_hc = [
            {"label": "Total Comprado Histórico", "value": f"${tot_comprado:,.2f}", "text_color": "1E40AF"},
            {"label": "Total Pagado / Abonado", "value": f"${tot_pagado:,.2f}", "text_color": "065F46"},
            {"label": "Saldo Pendiente Deudor", "value": f"${tot_saldo:,.2f}", "text_color": "991B1B"},
            {"label": "% Liquidado", "value": f"{pct_cub:.1f}%", "text_color": "065F46" if pct_cub >= 80 else "92400E"}
        ]
        next_r = agregar_seccion_kpis(ws, kpis_hc, start_row=4)

        headers_hc = ["Folio Factura", "Tipo Documento", "Emisión", "Vencimiento", "Días Crédito", "Días Retraso", "Estatus Pago", "Subtotal ($)", "Total Factura ($)", "Monto Pagado ($)", "Saldo Pendiente ($)"]
        rows_hc = [
            [
                f.get("folio"), f.get("tipo"), f.get("fecha_emision"), f.get("fecha_vencimiento"),
                f.get("dias_credito", 0), f.get("dias_retraso", 0), f.get("estatus_pago"),
                f.get("subtotal"), f.get("total"), f.get("monto_pagado"), f.get("saldo_pendiente")
            ]
            for f in facturas
        ]
        formatos_hc = {1: "center", 2: "badge", 3: "center", 4: "center", 5: "center", 6: "center", 7: "badge", 8: "currency", 9: "currency", 10: "currency", 11: "currency"}
        totales_hc = ["TOTALES", f"{len(facturas)} Facturas", "—", "—", "—", "—", "—", sum(float(f.get("subtotal", 0) or 0) for f in facturas), tot_comprado, tot_pagado, tot_saldo]

        agregar_tabla_ejecutiva(ws, next_r, headers_hc, rows_hc, formatos_hc, tema_color=COLOR_EMERALD_HEADER, totales_row=totales_hc)

        # HOJA 2: DESGLOSE DE PRODUCTOS VENDIDOS
        ws_prods = wb.create_sheet(title="Productos por Factura")
        aplicar_estilo_banner(
            ws_prods,
            titulo=f"DETALLE DE PRODUCTOS VENDIDOS — {cliente_nom.upper()}",
            subtitulo="Renglón por renglón de cada producto adquirido en las facturas",
            num_columnas=9,
            tema_color=COLOR_SLATE_DARK
        )

        headers_det = ["Folio Factura", "Fecha Emisión", "Código / SKU", "Descripción del Producto", "Cantidad", "Unidad", "Precio Unitario ($)", "Subtotal ($)", "Total con IVA ($)"]
        rows_det = []
        for f in facturas:
            fol = f.get("folio", "")
            f_em = f.get("fecha_emision", "")
            for p in f.get("productos", []):
                rows_det.append([
                    fol, f_em, p.get("sku") or "—", p.get("producto"),
                    p.get("cantidad"), p.get("uom") or "Pz", p.get("precio_unit"),
                    p.get("subtotal"), p.get("total")
                ])

        formatos_det = {1: "center", 2: "center", 3: "center", 4: "left", 5: "number", 6: "center", 7: "currency", 8: "currency", 9: "currency"}
        totales_det = ["TOTALES", "—", "—", f"{len(rows_det)} Renglones", sum(float(r[4] or 0) for r in rows_det), "—", "—", sum(float(r[7] or 0) for r in rows_det), sum(float(r[8] or 0) for r in rows_det)]
        agregar_tabla_ejecutiva(ws_prods, 4, headers_det, rows_det, formatos_det, tema_color=COLOR_SLATE_DARK, totales_row=totales_det)

    elif tipo == "comisiones":
        # ── REPORTE DE COMISIONES ─────────────────────────────────────────
        ws.title = "Comisiones Vendedores"
        tot_cob = sum(float(i.get("cobrado", 0) or 0) for i in items)
        tot_com = sum(float(i.get("comision", 0) or 0) for i in items)
        
        aplicar_estilo_banner(
            ws,
            titulo="CHALLENGER AUTOMOTIVE — LIQUIDACIÓN DE COMISIONES POR VENDEDOR",
            subtitulo="Desglose de cobros conciliados en Odoo y cálculo de comisión por producto",
            num_columnas=10,
            tema_color=COLOR_SLATE_DARK
        )
        
        kpis_com = [
            {"label": "Total Cobrado", "value": f"${tot_cob:,.2f}", "text_color": "1E40AF"},
            {"label": "Total Comisión a Pagar", "value": f"${tot_com:,.2f}", "text_color": "065F46"},
            {"label": "Pagos Conciliados", "value": f"{len(items)} pagos", "text_color": "581C87"}
        ]
        next_r = agregar_seccion_kpis(ws, kpis_com, start_row=4)
        
        headers_com = ["Fecha", "Cliente / Razón Social", "Factura(s)", "Pago #", "Monto Cobrado ($)", "Base Challenger ($)", "Base Multi/Pasta ($)", "Comisión Calculada ($)", "Estatus Tiempo", "Días Retraso"]
        rows_com = [
            [
                i.get("fecha"), i.get("cliente"), i.get("factura"),
                i.get("num_pago"), i.get("cobrado"), i.get("base_challenger"),
                i.get("base_multi"), i.get("comision"),
                "🟢 En tiempo" if i.get("en_tiempo") else f"⚠️ {i.get('dias_retraso', 0)}d retraso", i.get("dias_retraso", 0)
            ]
            for i in items
        ]
        formatos_com = {1: "center", 2: "left", 3: "left", 4: "center", 5: "currency", 6: "currency", 7: "currency", 8: "currency", 9: "badge", 10: "center"}
        totales_com = ["TOTALES", f"{len(items)} Pagos", "—", "—", tot_cob, sum(float(i.get("base_challenger", 0) or 0) for i in items), sum(float(i.get("base_multi", 0) or 0) for i in items), tot_com, "—", "—"]
        
        agregar_tabla_ejecutiva(ws, next_r, headers_com, rows_com, formatos_com, tema_color=COLOR_SLATE_DARK, totales_row=totales_com)

    elif tipo == "stock":
        # ── REPORTE DE STOCK ÓPTIMO & PLAN DE REABASTECIMIENTO ────────────
        ws.title = "Stock Óptimo y Reorden"
        tot_disp = sum(float(i.get("disponible", 0) or 0) for i in items)
        tot_val = sum(float(i.get("valor_estimado", 0) or 0) for i in items)
        tot_v_mes = sum(float(i.get("venta_mensual_est", 0) or 0) for i in items)
        tot_optimo = sum(float(i.get("stock_sugerido_mes", 0) or 0) for i in items)
        tot_pedir = sum(float(i.get("sugerido_pedir", 0) or 0) for i in items)
        tot_inv_reorden = sum(float(i.get("inversion_reorden", 0) or 0) for i in items)
        
        aplicar_estilo_banner(
            ws,
            titulo="CHALLENGER AUTOMOTIVE — PLAN DE STOCK ÓPTIMO Y SUGERENCIAS DE COMPRA",
            subtitulo="Estimación de demanda mensual (30d), stock mínimo de seguridad, stock meta óptimo y faltante a pedir",
            num_columnas=15,
            tema_color=COLOR_BLUE_HEADER
        )
        
        kpis_stk = [
            {"label": "Stock A la mano Total", "value": f"{int(tot_disp):,} uds", "text_color": "1E40AF"},
            {"label": "Demanda Mensual Est.", "value": f"{int(tot_v_mes):,} uds/mes", "text_color": "581C87"},
            {"label": "Stock Óptimo Meta (30d)", "value": f"{int(tot_optimo):,} uds", "text_color": "065F46"},
            {"label": "Total a Pedir (Faltante)", "value": f"{int(tot_pedir):,} uds", "text_color": "991B1B"},
            {"label": "Inversión Estimada Reorden", "value": f"${tot_inv_reorden:,.2f}", "text_color": "065F46"}
        ]
        next_r = agregar_seccion_kpis(ws, kpis_stk, start_row=4)
        
        headers_stk = [
            "Posición", "Producto / Código", "Almacén", "Categoría / Marca",
            "A la mano", "Reservado", "Disponible", "Demanda Mensual (30d)",
            "Stock Mínimo (15d)", "Stock Óptimo Meta (30d)", "Sugerido a Pedir",
            "Precio Venta + IVA ($)", "Inversión Reorden ($ MXN)", "Días Cobertura", "Estatus Stock"
        ]
        rows_stk = [
            [
                i.get("posicion"), i.get("producto"), i.get("almacen"),
                i.get("clasificacion"), i.get("stock_actual"), i.get("reservado"),
                i.get("disponible"), i.get("venta_mensual_est"),
                i.get("stock_minimo"), i.get("stock_sugerido_mes"),
                i.get("sugerido_pedir"),
                i.get("precio_unitario_est", 0.0),
                i.get("inversion_reorden", 0.0),
                f"{i.get('dias_cobertura', 0):.1f}d" if (i.get("dias_cobertura", 0) or 0) < 900 else "∞",
                i.get("estado")
            ]
            for i in items
        ]
        formatos_stk = {
            1: "center", 2: "left", 3: "left", 4: "badge",
            5: "number", 6: "number", 7: "number", 8: "number",
            9: "number", 10: "number", 11: "number",
            12: "currency", 13: "currency", 14: "center", 15: "badge"
        }
        totales_stk = [
            "TOTALES", f"{len(items)} Productos", "—", "—",
            sum(float(i.get("stock_actual", 0) or 0) for i in items),
            sum(float(i.get("reservado", 0) or 0) for i in items),
            tot_disp, tot_v_mes,
            sum(float(i.get("stock_minimo", 0) or 0) for i in items),
            tot_optimo, tot_pedir, "—",
            tot_inv_reorden, "—", "—"
        ]
        
        agregar_tabla_ejecutiva(ws, next_r, headers_stk, rows_stk, formatos_stk, tema_color=COLOR_BLUE_HEADER, totales_row=totales_stk)

        # ── HOJA 2: RESUMEN CONSOLIDADO POR ALMACÉN ────────────────────────
        # Agrupar métricas por almacén
        alms_map = {}
        for i in items:
            alm = i.get("almacen") or "Sin Almacén"
            if alm not in alms_map:
                alms_map[alm] = {
                    "almacen": alm, "productos": 0, "disponible": 0.0,
                    "ventas_mes": 0.0, "stock_optimo": 0.0, "sugerido_pedir": 0.0,
                    "inversion_reorden": 0.0, "agotados": 0, "reordenar": 0, "optimos": 0
                }
            m = alms_map[alm]
            m["productos"] += 1
            m["disponible"] += float(i.get("disponible", 0) or 0)
            m["ventas_mes"] += float(i.get("venta_mensual_est", 0) or 0)
            m["stock_optimo"] += float(i.get("stock_sugerido_mes", 0) or 0)
            m["sugerido_pedir"] += float(i.get("sugerido_pedir", 0) or 0)
            m["inversion_reorden"] += float(i.get("inversion_reorden", 0) or 0)
            st = i.get("estado")
            if st == "AGOTADO": m["agotados"] += 1
            elif st == "REORDENAR": m["reordenar"] += 1
            elif st == "OPTIMO": m["optimos"] += 1

        ws_alm = wb.create_sheet(title="Resumen por Almacén")
        aplicar_estilo_banner(
            ws_alm,
            titulo="CHALLENGER AUTOMOTIVE — RESUMEN DE STOCK ÓPTIMO POR ALMACÉN",
            subtitulo="Balance de inventario, demanda mensual proyectada y presupuesto de compra por ubicación",
            num_columnas=10,
            tema_color=COLOR_EMERALD_HEADER
        )

        headers_alm = [
            "Almacén / Ubicación", "Total Productos", "Stock Disponible",
            "Demanda Mensual Est.", "Stock Óptimo Meta", "Total a Pedir",
            "Inversión Reorden ($ MXN)", "🚨 Agotados", "🟡 Bajo Mínimo", "🟢 Óptimos"
        ]
        rows_alm = [
            [
                a["almacen"], a["productos"], a["disponible"],
                a["ventas_mes"], a["stock_optimo"], a["sugerido_pedir"],
                a["inversion_reorden"], a["agotados"], a["reordenar"], a["optimos"]
            ]
            for a in sorted(alms_map.values(), key=lambda x: x["inversion_reorden"], reverse=True)
        ]
        formatos_alm = {
            1: "left", 2: "number", 3: "number", 4: "number",
            5: "number", 6: "number", 7: "currency", 8: "number", 9: "number", 10: "number"
        }
        totales_alm = [
            "TOTALES", len(items), tot_disp, tot_v_mes, tot_optimo, tot_pedir,
            tot_inv_reorden,
            sum(a["agotados"] for a in alms_map.values()),
            sum(a["reordenar"] for a in alms_map.values()),
            sum(a["optimos"] for a in alms_map.values())
        ]
        agregar_tabla_ejecutiva(ws_alm, 4, headers_alm, rows_alm, formatos_alm, tema_color=COLOR_EMERALD_HEADER, totales_row=totales_alm)

    elif tipo == "sat":
        # ── REPORTE SAT PASIVOS ───────────────────────────────────────────
        ws.title = "Pasivos Proveedores SAT"
        tot_sat = sum(float(i.get("saldo_pendiente", 0) or 0) for i in items)
        
        aplicar_estilo_banner(
            ws,
            titulo="CHALLENGER AUTOMOTIVE — CUENTAS POR PAGAR Y FACTURAS SAT (PROVEEDORES)",
            subtitulo="Conciliación fiscal, pasivos pendientes y seguimiento de Complementos de Pago (REP)",
            num_columnas=11,
            tema_color=COLOR_AMBER_HEADER
        )
        
        headers_sat = ["Posición", "Proveedor / Razón Social", "RFC Proveedor", "UUID / Folio Fiscal", "Fecha Factura", "Método", "Total Factura ($)", "Monto Abonado ($)", "Saldo Pendiente ($)", "Estatus Pago", "Estatus REP"]
        rows_sat = [
            [
                i.get("posicion"), i.get("proveedor"), i.get("rfc_proveedor"),
                i.get("uuid"), i.get("fecha_factura"), i.get("metodo_pago"),
                i.get("total_factura"), i.get("monto_abonado"), i.get("saldo_pendiente"),
                i.get("estatus_pago"), i.get("estatus_rep")
            ]
            for i in items
        ]
        formatos_sat = {1: "center", 2: "left", 3: "center", 4: "center", 5: "center", 6: "center", 7: "currency", 8: "currency", 9: "currency", 10: "badge", 11: "badge"}
        totales_sat = ["TOTALES", f"{len(items)} Facturas", "—", "—", "—", "—", sum(float(i.get("total_factura", 0) or 0) for i in items), sum(float(i.get("monto_abonado", 0) or 0) for i in items), tot_sat, "—", "—"]
        
        agregar_tabla_ejecutiva(ws, 4, headers_sat, rows_sat, formatos_sat, tema_color=COLOR_AMBER_HEADER, totales_row=totales_sat)

    elif tipo == "ventas_vendedor":
        # ── REPORTE DE VENTAS POR VENDEDOR & PRODUCTOS VENDIDOS ───────────
        vv_data = data.get("vv_data") or {}
        resumen_vends = vv_data.get("resumen_vendedores") or items or []
        totales = vv_data.get("totales_generales") or {}
        periodo_nom = vv_data.get("periodo_seleccionado", "Mes Actual")

        # HOJA 1: RESUMEN POR VENDEDOR
        ws.title = "Ventas por Vendedor"
        aplicar_estilo_banner(
            ws,
            titulo=f"CHALLENGER AUTOMOTIVE — VENTAS POR VENDEDOR ({periodo_nom.upper()})",
            subtitulo="Facturación consolidada por vendedor oficial, volumen y participación de mercado",
            num_columnas=9,
            tema_color=COLOR_EMERALD_HEADER
        )

        tot_fact = totales.get("total_facturado", sum(float(v.get("total_subtotal", 0) or 0) for v in resumen_vends))
        tot_units = totales.get("total_unidades", sum(float(v.get("total_cantidad", 0) or 0) for v in resumen_vends))
        v_lider = totales.get("vendedor_lider", "—")

        kpis_vv = [
            {"label": f"Total Facturado ({periodo_nom})", "value": f"${tot_fact:,.2f}", "text_color": "065F46"},
            {"label": "Unidades Vendidas", "value": f"{int(tot_units):,} uds", "text_color": "1E40AF"},
            {"label": "Vendedor Líder", "value": f"{v_lider}", "text_color": "581C87"},
            {"label": "Vendedores Activos", "value": f"{len(resumen_vends)}", "text_color": "92400E"}
        ]
        next_r_vv = agregar_seccion_kpis(ws, kpis_vv, start_row=4)

        headers_vv = ["Posición", "Vendedor Oficial", "Total Facturado ($)", "Unidades Vendidas", "Facturas #", "Clientes #", "Ticket Promedio ($)", "% Participación", "Top Producto"]
        rows_vv = [
            [
                v.get("posicion"), v.get("vendedor"), v.get("total_subtotal"),
                v.get("total_cantidad"), v.get("num_facturas"), v.get("num_clientes"),
                v.get("ticket_promedio"), v.get("pct_participacion"), v.get("top_producto")
            ]
            for v in resumen_vends
        ]
        formatos_vv = {1: "center", 2: "left", 3: "currency", 4: "number", 5: "center", 6: "center", 7: "currency", 8: "percent", 9: "left"}
        totales_vv = ["TOTALES", f"{len(resumen_vends)} Vendedores", tot_fact, tot_units, sum(int(v.get("num_facturas", 0) or 0) for v in resumen_vends), "—", "—", "100.0%", "—"]
        agregar_tabla_ejecutiva(ws, next_r_vv, headers_vv, rows_vv, formatos_vv, tema_color=COLOR_EMERALD_HEADER, totales_row=totales_vv)

        # HOJA 2: DESGLOSE DE PRODUCTOS POR VENDEDOR
        ws_prods = wb.create_sheet(title="Productos por Vendedor")
        aplicar_estilo_banner(
            ws_prods,
            titulo=f"CHALLENGER AUTOMOTIVE — DESGLOSE DE PRODUCTOS POR VENDEDOR ({periodo_nom.upper()})",
            subtitulo="Detalle de productos comercializados por cada ejecutivo de ventas",
            num_columnas=9,
            tema_color=COLOR_SLATE_DARK
        )

        headers_pv = ["Pos", "Vendedor", "Producto / Código", "Categoría / Marca", "Cantidad (Uds)", "Precio Prom ($)", "Total Facturado ($)", "% del Vendedor", "Facturas #"]
        rows_pv = []
        for v in resumen_vends:
            v_name = v.get("vendedor", "Vendedor")
            for p in v.get("productos_vendidos", []):
                rows_pv.append([
                    p.get("posicion"), v_name, p.get("producto"),
                    p.get("categoria"), p.get("cantidad"), p.get("precio_promedio"),
                    p.get("subtotal"), p.get("pct_del_vendedor"), p.get("num_facturas")
                ])
        formatos_pv = {1: "center", 2: "left", 3: "left", 4: "badge", 5: "number", 6: "currency", 7: "currency", 8: "percent", 9: "center"}
        totales_pv = ["TOTALES", "—", f"{len(rows_pv)} Líneas", "—", sum(float(r[4] or 0) for r in rows_pv), "—", sum(float(r[6] or 0) for r in rows_pv), "—", "—"]
        agregar_tabla_ejecutiva(ws_prods, 4, headers_pv, rows_pv, formatos_pv, tema_color=COLOR_SLATE_DARK, totales_row=totales_pv)

        # HOJA 3: MATRIZ MULTI-MES POR VENDEDOR
        matriz_m = vv_data.get("matriz_meses_vendedores", [])
        periodos_h = vv_data.get("periodos_historico", [])
        if matriz_m and periodos_h:
            ws_m_vend = wb.create_sheet(title="Histórico por Vendedor")
            aplicar_estilo_banner(
                ws_m_vend,
                titulo="CHALLENGER AUTOMOTIVE — EVOLUCIÓN HISTÓRICA POR VENDEDOR (ÚLTIMOS 12 MESES)",
                subtitulo=f"Comportamiento de facturación mensual por ejecutivo",
                num_columnas=len(periodos_h) + 4,
                tema_color=COLOR_BLUE_HEADER
            )
            headers_mv = ["Pos", "Vendedor Oficial"] + [f"{per} ($)" for per in periodos_h] + ["Total Histórico ($)", "Crecimiento MoM (%)"]
            rows_mv = []
            for mv in matriz_m:
                r_row = [mv.get("posicion"), mv.get("vendedor")]
                for per in periodos_h:
                    r_row.append(mv.get("ventas_por_mes", {}).get(per, 0.0))
                r_row.append(mv.get("total_historico"))
                r_row.append(mv.get("mom_growth"))
                rows_mv.append(r_row)
            formatos_mv = {1: "center", 2: "left"}
            for c_i in range(3, len(periodos_h) + 4):
                formatos_mv[c_i] = "currency"
            formatos_mv[len(periodos_h) + 4] = "percent"
            agregar_tabla_ejecutiva(ws_m_vend, 4, headers_mv, rows_mv, formatos_mv, tema_color=COLOR_BLUE_HEADER)

    elif tipo == "ventas_vendedor_3meses":
        # ── REPORTE TRIMESTRAL (3 MESES) POR VENDEDOR & PRODUCTOS ─────────
        vv_data = data.get("vv_data") or {}
        resumen_3m = vv_data.get("resumen_vendedores_3m") or items or []
        totales = vv_data.get("totales_generales_3m") or {}
        periodos_nombres = vv_data.get("periodos_meses_nombres") or ["Mes 1", "Mes 2", "Mes 3"]
        periodo_label = vv_data.get("periodo_trimestre", "Últimos 3 Meses")

        # HOJA 1: RESUMEN TRIMESTRAL POR VENDEDOR
        ws.title = "Resumen 3 Meses"
        aplicar_estilo_banner(
            ws,
            titulo=f"CHALLENGER AUTOMOTIVE — REPORTE DE VENTAS POR VENDEDOR ({periodo_label.upper()})",
            subtitulo=f"Facturación comparativa de los últimos 3 meses por ejecutivo y tendencias de crecimiento",
            num_columnas=11,
            tema_color=COLOR_PURPLE_HEADER
        )

        tot_fact = totales.get("total_facturado", sum(float(v.get("total_subtotal_3m", 0) or 0) for v in resumen_3m))
        tot_units = totales.get("total_unidades", sum(float(v.get("total_cantidad_3m", 0) or 0) for v in resumen_3m))
        v_lider = totales.get("vendedor_lider", "—")

        kpis_3m = [
            {"label": f"Facturación Trimestral ({periodo_label})", "value": f"${tot_fact:,.2f}", "text_color": "581C87"},
            {"label": "Unidades Totales (3 Meses)", "value": f"{int(tot_units):,} uds", "text_color": "1E40AF"},
            {"label": "Vendedor Líder Trimestral", "value": f"{v_lider}", "text_color": "065F46"},
            {"label": "Ejecutivos Evaluados", "value": f"{len(resumen_3m)}", "text_color": "92400E"}
        ]
        next_r_3m = agregar_seccion_kpis(ws, kpis_3m, start_row=4)

        headers_3m = [
            "Posición", "Vendedor Oficial",
            f"{periodos_nombres[0]} ($)", f"{periodos_nombres[1]} ($)", f"{periodos_nombres[2]} ($)",
            "Total 3 Meses ($)", "Unidades (3M)", "Facturas #", "Clientes #", "% Participación", "Tendencia MoM"
        ]
        rows_3m = [
            [
                v.get("posicion"), v.get("vendedor"),
                v.get("sub_mes_1"), v.get("sub_mes_2"), v.get("sub_mes_3"),
                v.get("total_subtotal_3m"), v.get("total_cantidad_3m"),
                v.get("num_facturas_3m"), v.get("num_clientes_3m"),
                v.get("pct_participacion_3m"), v.get("tendencia_mom")
            ]
            for v in resumen_3m
        ]
        formatos_3m = {
            1: "center", 2: "left", 3: "currency", 4: "currency", 5: "currency",
            6: "currency", 7: "number", 8: "center", 9: "center", 10: "percent", 11: "percent"
        }
        totales_3m = [
            "TOTALES", f"{len(resumen_3m)} Vendedores",
            sum(float(v.get("sub_mes_1", 0) or 0) for v in resumen_3m),
            sum(float(v.get("sub_mes_2", 0) or 0) for v in resumen_3m),
            sum(float(v.get("sub_mes_3", 0) or 0) for v in resumen_3m),
            tot_fact, tot_units,
            sum(int(v.get("num_facturas_3m", 0) or 0) for v in resumen_3m),
            "—", "100.0%", "—"
        ]
        agregar_tabla_ejecutiva(ws, next_r_3m, headers_3m, rows_3m, formatos_3m, tema_color=COLOR_PURPLE_HEADER, totales_row=totales_3m)

        # HOJA 2: DESGLOSE DE PRODUCTOS POR VENDEDOR EN 3 MESES
        ws_prods3 = wb.create_sheet(title="Productos 3 Meses")
        aplicar_estilo_banner(
            ws_prods3,
            titulo=f"CHALLENGER AUTOMOTIVE — PRODUCTOS POR VENDEDOR EN LOS ÚLTIMOS 3 MESES",
            subtitulo=f"Detalle mes a mes de cada producto comercializado ({periodo_label})",
            num_columnas=12,
            tema_color=COLOR_SLATE_DARK
        )

        headers_p3 = [
            "Pos", "Vendedor", "Producto / Código", "Categoría / Marca",
            f"Cant {periodos_nombres[0]}", f"Cant {periodos_nombres[1]}", f"Cant {periodos_nombres[2]}",
            "Cant Total (3M)", "P. Prom ($)", "Total Facturado 3M ($)", "% del Vendedor", "Facturas #"
        ]
        rows_p3 = []
        for v in resumen_3m:
            v_name = v.get("vendedor", "Vendedor")
            for p in v.get("productos_vendidos_3m", []):
                rows_p3.append([
                    p.get("posicion"), v_name, p.get("producto"), p.get("categoria"),
                    p.get("cant_mes_1"), p.get("cant_mes_2"), p.get("cant_mes_3"),
                    p.get("total_cantidad_3m"), p.get("precio_promedio_3m"),
                    p.get("total_subtotal_3m"), p.get("pct_del_vendedor_3m"), p.get("num_facturas_3m")
                ])
        formatos_p3 = {
            1: "center", 2: "left", 3: "left", 4: "badge",
            5: "number", 6: "number", 7: "number", 8: "number",
            9: "currency", 10: "currency", 11: "percent", 12: "center"
        }
        totales_p3 = [
            "TOTALES", "—", f"{len(rows_p3)} Registros", "—",
            sum(float(r[4] or 0) for r in rows_p3), sum(float(r[5] or 0) for r in rows_p3), sum(float(r[6] or 0) for r in rows_p3),
            sum(float(r[7] or 0) for r in rows_p3), "—", sum(float(r[9] or 0) for r in rows_p3), "—", "—"
        ]
        agregar_tabla_ejecutiva(ws_prods3, 4, headers_p3, rows_p3, formatos_p3, tema_color=COLOR_SLATE_DARK, totales_row=totales_p3)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(tmp.name)
    tmp.close()

    return send_file(tmp.name, as_attachment=True, download_name=f"Reporte_Ejecutivo_{tipo}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx")


@app.route("/api/export_pdf", methods=["POST"])
def export_pdf():
    data = request.get_json() or {}
    tipo = data.get("tipo", "productos")
    items = data.get("items", [])
    ml_data = data.get("ml_data")
    hist_data = data.get("hist_data")
    cobrado_mes_data = data.get("cobrado_mes_data")
    periodo_filtro = data.get("periodo_filtro", "TODOS")
    vendedor_filtro = data.get("vendedor_filtro", "TODOS")
    empresa_nombre = data.get("empresa_nombre", "TODAS (Consolidado)")

    try:
        from pdf_styler import generar_pdf_reporte
        pdf_path = generar_pdf_reporte(
            tipo=tipo,
            items=items,
            ml_data=ml_data,
            hist_data=hist_data,
            cobrado_mes_data=cobrado_mes_data,
            periodo_filtro=periodo_filtro,
            vendedor_filtro=vendedor_filtro,
            empresa_nombre=empresa_nombre
        )
        
        if tipo in ["historial_cliente", "compras_cliente"] and hist_data:
            c_nom = str(hist_data.get("cliente_nombre", "Cliente")).replace(" ", "_").replace("/", "_")
            d_name = f"Estado_Cuenta_{c_nom}_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
        elif tipo == "cobrado_mes":
            d_name = f"Reporte_Cobrado_Mes_{periodo_filtro}_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
        else:
            d_name = f"Reporte_Ejecutivo_{tipo}_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
            
        return send_file(
            pdf_path,
            as_attachment=True,
            download_name=d_name,
            mimetype="application/pdf"
        )
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": f"Error generando PDF: {str(e)}"})


# ── MÓDULO DE GASTOS A PROVEEDORES Y CONCILIACIÓN BANCARIA ───────────────────
from gastos_conciliacion_engine import (
    ejecutar_conciliacion_engine, cargar_json, guardar_json,
    COMPROBANTES_FILE, ESTADOS_CUENTA_FILE, CONCILIACIONES_FILE,
    COMPROBANTES_DIR, ESTADOS_CUENTA_DIR,
    extraer_datos_pdf_comprobante,
    parsear_estado_cuenta_archivo, parsear_texto_movimientos_bancarios,
    generar_excel_conciliacion, parsear_fecha_flexible
)

@app.route("/api/gastos/conciliacion", methods=["GET", "POST"])
def api_gastos_conciliacion():
    if request.method == "POST":
        data = request.get_json() or {}
        empresa = data.get("empresa", "TODAS (Consolidado)")
        mes = data.get("mes")
        anio = data.get("anio")
        banco = data.get("banco", "TODOS")
    else:
        empresa = request.args.get("empresa", "TODAS (Consolidado)")
        mes = request.args.get("mes")
        anio = request.args.get("anio")
        banco = request.args.get("banco", "TODOS")

    try:
        mes_val = int(mes) if mes and str(mes).isdigit() else None
        anio_val = int(anio) if anio and str(anio).isdigit() else None
        res = ejecutar_conciliacion_engine(empresa=empresa, mes=mes_val, anio=anio_val, banco_filtro=banco)
        return jsonify({"success": True, "data": res})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)})

@app.route("/api/gastos/comprobantes", methods=["GET"])
def api_gastos_listar_comprobantes():
    empresa = request.args.get("empresa")
    comprobantes = cargar_json(COMPROBANTES_FILE, [])
    if empresa and empresa != "TODAS (Consolidado)":
        comprobantes = [c for c in comprobantes if c.get("empresa") == empresa]
    return jsonify({"success": True, "comprobantes": comprobantes})

@app.route("/api/gastos/comprobante/upload", methods=["POST"])
def api_gastos_subir_comprobante():
    try:
        empresa = request.form.get("empresa", "BALDOMERO MEDINA GARZA")
        nota_whatsapp = request.form.get("nota_whatsapp", "").strip()
        proveedor_asignado = (request.form.get("proveedor_asignado") or request.form.get("proveedor") or "").strip() or nota_whatsapp
        fecha_manual = request.form.get("fecha", "").strip()
        importe_manual = (request.form.get("importe") or request.form.get("monto") or "").strip()
        banco_origen = (request.form.get("banco_origen") or request.form.get("banco") or "").strip()
        concepto_manual = request.form.get("concepto", "").strip()
        referencia_manual = request.form.get("referencia", "").strip()
        clave_rastreo_manual = (request.form.get("clave_rastreo") or request.form.get("rastreo") or "").strip()
        comision_manual = (request.form.get("comision") or "0").strip()

        archivo = request.files.get("archivo") or request.files.get("file")
        archivo_nombre = ""
        archivo_ruta = ""
        datos_extraidos = {}

        nuevo_id = f"COMP-{datetime.now().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:6]}"

        if archivo and archivo.filename:
            archivo_nombre = archivo.filename
            ext = os.path.splitext(archivo_nombre)[1].lower()
            save_name = f"{nuevo_id}_{archivo_nombre}"
            archivo_ruta = os.path.join(COMPROBANTES_DIR, save_name)
            archivo.save(archivo_ruta)

            if ext == ".pdf":
                datos_extraidos = extraer_datos_pdf_comprobante(archivo_ruta)

        fecha_final = fecha_manual or datos_extraidos.get("fecha") or datetime.now().strftime("%Y-%m-%d")
        fecha_final = parsear_fecha_flexible(fecha_final)

        try:
            importe_final = float(importe_manual) if importe_manual else float(datos_extraidos.get("importe", 0.0))
        except Exception:
            importe_final = float(datos_extraidos.get("importe", 0.0))

        try:
            comision_final = float(comision_manual) if comision_manual else float(datos_extraidos.get("comision", 0.0))
        except Exception:
            comision_final = float(datos_extraidos.get("comision", 0.0))

        total_final = round(importe_final + comision_final, 2)

        proveedor_final = proveedor_asignado or datos_extraidos.get("beneficiario") or nota_whatsapp or "Proveedor Varios"
        banco_final = banco_origen or datos_extraidos.get("banco_origen") or "BBVA"
        concepto_final = concepto_manual or datos_extraidos.get("concepto") or nota_whatsapp or "Pago Proveedor"
        rastreo_final = clave_rastreo_manual or datos_extraidos.get("clave_rastreo") or ""
        ref_final = referencia_manual or datos_extraidos.get("referencia") or ""

        comprobante_obj = {
            "id": nuevo_id,
            "empresa": empresa,
            "fecha": fecha_final,
            "hora": datos_extraidos.get("hora", datetime.now().strftime("%H:%M:%S")),
            "banco_origen": banco_final,
            "cuenta_origen": datos_extraidos.get("cuenta_origen", "Principal"),
            "banco_destino": datos_extraidos.get("banco_destino", "Destino"),
            "cuenta_destino": datos_extraidos.get("cuenta_destino", ""),
            "beneficiario": datos_extraidos.get("beneficiario") or proveedor_final,
            "proveedor_asignado": proveedor_final,
            "proveedor": proveedor_final,
            "concepto": concepto_final,
            "nota_whatsapp": nota_whatsapp,
            "referencia": ref_final,
            "clave_rastreo": rastreo_final,
            "folio_operacion": datos_extraidos.get("folio_operacion", ""),
            "importe": importe_final,
            "monto": importe_final,
            "comision": comision_final,
            "total": total_final,
            "monto_total": total_final,
            "tipo_comprobante": datos_extraidos.get("tipo_comprobante", "Comprobante"),
            "archivo_nombre": archivo_nombre,
            "archivo_ruta": archivo_ruta,
            "creado_en": datetime.now().isoformat()
        }

        comprobantes = cargar_json(COMPROBANTES_FILE, [])
        comprobantes.insert(0, comprobante_obj)
        guardar_json(COMPROBANTES_FILE, comprobantes)

        return jsonify({"success": True, "comprobante": comprobante_obj, "message": "Comprobante subido y procesado con éxito."})
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)})

@app.route("/api/gastos/comprobante/<comp_id>", methods=["DELETE", "PUT"])
def api_gastos_comprobante_ops(comp_id):
    comprobantes = cargar_json(COMPROBANTES_FILE, [])
    if request.method == "DELETE":
        comprobantes = [c for c in comprobantes if c["id"] != comp_id]
        guardar_json(COMPROBANTES_FILE, comprobantes)
        conciliaciones = cargar_json(CONCILIACIONES_FILE, {})
        conciliaciones = {m: c for m, c in conciliaciones.items() if c != comp_id}
        guardar_json(CONCILIACIONES_FILE, conciliaciones)
        return jsonify({"success": True, "message": "Comprobante eliminado con éxito."})
    
    if request.method == "PUT":
        data = request.get_json() or {}
        for c in comprobantes:
            if c["id"] == comp_id:
                for k, v in data.items():
                    if k in c:
                        c[k] = v
                break
        guardar_json(COMPROBANTES_FILE, comprobantes)
        return jsonify({"success": True, "message": "Comprobante actualizado con éxito."})

@app.route("/api/gastos/banco/upload", methods=["POST"])
def api_gastos_subir_estado_cuenta():
    try:
        empresa = request.form.get("empresa", "BALDOMERO MEDINA GARZA")
        banco = request.form.get("banco", "BBVA")
        cuenta_nombre = request.form.get("cuenta_nombre", "").strip() or f"{banco} Principal"
        texto_pegado = (request.form.get("texto_pegado") or request.form.get("raw_text") or "").strip()

        archivo = request.files.get("archivo") or request.files.get("file")
        movimientos_procesados = []

        if archivo and archivo.filename:
            archivo_nombre = archivo.filename
            save_name = f"EDO_{datetime.now().strftime('%Y%m%d%H%M%S')}_{archivo_nombre}"
            archivo_ruta = os.path.join(ESTADOS_CUENTA_DIR, save_name)
            archivo.save(archivo_ruta)
            movimientos_procesados = parsear_estado_cuenta_archivo(archivo_ruta, banco, empresa, cuenta_nombre)
        elif texto_pegado:
            movimientos_procesados = parsear_texto_movimientos_bancarios(texto_pegado, banco, empresa, cuenta_nombre)
            if movimientos_procesados:
                existentes = cargar_json(ESTADOS_CUENTA_FILE, [])
                existentes.extend(movimientos_procesados)
                guardar_json(ESTADOS_CUENTA_FILE, existentes)

        return jsonify({
            "success": True,
            "movimientos": movimientos_procesados,
            "conteo": len(movimientos_procesados),
            "message": f"Se procesaron {len(movimientos_procesados)} movimientos bancarios correctamente."
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"success": False, "error": str(e)})

@app.route("/api/gastos/banco/movimientos", methods=["GET"])
def api_gastos_listar_movimientos():
    empresa = request.args.get("empresa")
    movimientos = cargar_json(ESTADOS_CUENTA_FILE, [])
    if empresa and empresa != "TODAS (Consolidado)":
        movimientos = [m for m in movimientos if m.get("empresa") == empresa]
    return jsonify({"success": True, "movimientos": movimientos})

@app.route("/api/gastos/banco/movimiento/<mov_id>", methods=["DELETE"])
def api_gastos_eliminar_movimiento(mov_id):
    movimientos = cargar_json(ESTADOS_CUENTA_FILE, [])
    movimientos = [m for m in movimientos if m["id"] != mov_id]
    guardar_json(ESTADOS_CUENTA_FILE, movimientos)
    conciliaciones = cargar_json(CONCILIACIONES_FILE, {})
    if mov_id in conciliaciones:
        del conciliaciones[mov_id]
        guardar_json(CONCILIACIONES_FILE, conciliaciones)
    return jsonify({"success": True, "message": "Movimiento eliminado."})

@app.route("/api/gastos/conciliar/vincular", methods=["POST"])
def api_gastos_vincular_manual():
    data = request.get_json() or {}
    mov_id = data.get("movimiento_id")
    comp_id = data.get("comprobante_id")
    desvincular = data.get("desvincular", False)

    conciliaciones = cargar_json(CONCILIACIONES_FILE, {})
    if desvincular:
        if mov_id in conciliaciones:
            del conciliaciones[mov_id]
            guardar_json(CONCILIACIONES_FILE, conciliaciones)
        return jsonify({"success": True, "message": "Movimiento desvinculado con éxito."})

    if not mov_id or not comp_id:
        return jsonify({"success": False, "error": "Faltan IDs para vincular."})

    conciliaciones[mov_id] = comp_id
    guardar_json(CONCILIACIONES_FILE, conciliaciones)
    return jsonify({"success": True, "message": "Conciliación manual vinculada exitosamente."})

@app.route("/api/gastos/comprobante/archivo/<comp_id>", methods=["GET"])
def api_gastos_ver_archivo(comp_id):
    comprobantes = cargar_json(COMPROBANTES_FILE, [])
    comp = next((c for c in comprobantes if c["id"] == comp_id), None)
    if not comp or not comp.get("archivo_ruta") or not os.path.exists(comp["archivo_ruta"]):
        return jsonify({"error": "Archivo no encontrado"}), 404
    
    return send_file(comp["archivo_ruta"])

@app.route("/api/gastos/exportar/excel", methods=["GET", "POST"])
def api_gastos_exportar_excel():
    empresa = request.args.get("empresa", "TODAS (Consolidado)")
    mes = request.args.get("mes")
    anio = request.args.get("anio")
    try:
        excel_path = generar_excel_conciliacion(empresa, mes, anio)
        return send_file(
            excel_path,
            as_attachment=True,
            download_name=f"Reporte_Conciliacion_Gastos_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5050))
    app.run(host="0.0.0.0", port=port, debug=False)

